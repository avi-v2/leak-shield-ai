"""
Modular Document Processing and Content Security Scanner for Leak-Shield-AI.

Extracts and inspects the actual contents of uploaded documents:
- Plaintext (.txt, .md)
- Structured Data (.csv, .json)
- Portable Document Format (.pdf) via pypdf
- Microsoft Word (.docx) via python-docx

Applies strict security controls:
1. File signature & extension validation (prevent executable spoofing).
2. Size boundary enforcement (prevent DoS memory exhaustion).
3. No code/macro execution.
4. Chunking for long documents with source location tracking (page, row, paragraph).
5. Comprehensive DLP and Prompt Injection detection reusing existing security engines.
"""

import io
import re
import csv
import json
import uuid
from datetime import datetime, timezone
from typing import List, Dict, Any, Tuple, Optional
from pathlib import Path

import pypdf
import docx

from app.config import settings
from app.schemas import Finding, FileScanResponse
from app.database import load_policies
from app.security.sensitive_data import scan_sensitive_data
from app.security.injection_detector import detect_prompt_injection
from app.security.policy_engine import evaluate_security_policy
from app.security.ocr_utils import is_ocr_available, extract_text_from_image

MAX_CHUNK_SIZE = 3500
CHUNK_OVERLAP = 300

# Executable / dangerous binary magic bytes to instantly reject
DISALLOWED_MAGIC_SIGNATURES = [
    (b"MZ", "Windows Executable/DLL (PE)"),
    (b"\x7fELF", "Linux ELF Executable"),
    (b"\xca\xfe\xba\xbe", "Java Class / Mach-O Binary"),
    (b"\xfe\xed\xfa\xce", "Mach-O Binary"),
    (b"\xfe\xed\xfa\xcf", "Mach-O Binary"),
    (b"#!", "Executable Shell Script")
]

def _detect_file_extension(filename: str) -> str:
    """Extract lowercase file extension including dot."""
    ext = Path(filename).suffix.lower()
    return ext

def _validate_file_signature(file_bytes: bytes, extension: str) -> Tuple[bool, Optional[str]]:
    """
    Validate file headers against dangerous magic bytes and verify format integrity.
    """
    if not file_bytes:
        return False, "File is empty (0 bytes)."

    # Reject dangerous executable formats spoofing extensions
    for magic, name in DISALLOWED_MAGIC_SIGNATURES:
        if file_bytes.startswith(magic):
            return False, f"Dangerous executable header detected ({name}). Upload blocked."

    # Validate specific format signatures
    if extension == ".pdf":
        if not file_bytes.startswith(b"%PDF"):
            return False, "Invalid PDF header: Missing '%PDF' signature."
    elif extension == ".docx":
        # DOCX is a zip container starting with PK\x03\x04
        if not file_bytes.startswith(b"PK\x03\x04"):
            return False, "Invalid DOCX header: File is not a valid OpenXML package."

    return True, None

def _decode_text_bytes(file_bytes: bytes) -> str:
    """Robustly decode raw text bytes trying common encodings."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")

def _chunk_text_with_locations(
    text: str,
    base_location: str
) -> List[Tuple[str, str]]:
    """
    Split text into overlapping chunks if it exceeds MAX_CHUNK_SIZE.
    Returns list of (chunk_text, location_descriptor).
    """
    if len(text) <= MAX_CHUNK_SIZE:
        return [(text, base_location)]

    chunks = []
    start = 0
    chunk_idx = 1
    total_len = len(text)

    while start < total_len:
        end = min(start + MAX_CHUNK_SIZE, total_len)
        chunk_str = text[start:end]
        loc = f"{base_location} (Chunk {chunk_idx}, char {start}-{end})"
        chunks.append((chunk_str, loc))
        chunk_idx += 1
        if end >= total_len:
            break
        start += MAX_CHUNK_SIZE - CHUNK_OVERLAP

    return chunks

def _extract_from_txt_or_md(file_bytes: bytes, filename: str) -> Tuple[List[Tuple[str, str]], str, int]:
    """Extract readable text from plain text or Markdown files."""
    text = _decode_text_bytes(file_bytes)
    chunks = _chunk_text_with_locations(text, "Document Text")
    return chunks, "SUCCESS", 1

def _extract_from_csv(file_bytes: bytes, filename: str) -> Tuple[List[Tuple[str, str]], str, int]:
    """Extract and structure data from CSV files with row and column tracking."""
    text = _decode_text_bytes(file_bytes)
    reader = csv.reader(io.StringIO(text))
    chunks: List[Tuple[str, str]] = []
    
    header = None
    row_count = 0

    for idx, row in enumerate(reader, start=1):
        if not row:
            continue
        if idx == 1:
            header = row
            # Scan header line
            header_text = " , ".join(header)
            chunks.append((header_text, "CSV Header Row"))
            continue

        row_count += 1
        # Include column names with values for contextual secret matching
        formatted_row = []
        for col_idx, val in enumerate(row):
            col_name = header[col_idx] if header and col_idx < len(header) else f"Col{col_idx+1}"
            formatted_row.append(f"{col_name}: {val}")
        
        row_str = " | ".join(formatted_row)
        chunks.append((row_str, f"CSV Row {idx}"))

    return chunks, "SUCCESS", row_count

def _extract_from_json(file_bytes: bytes, filename: str) -> Tuple[List[Tuple[str, str]], str, int]:
    """Extract values from JSON documents with key and object path tracking."""
    text = _decode_text_bytes(file_bytes)
    try:
        data = json.loads(text)
    except Exception as e:
        # Fall back to raw text if invalid JSON
        chunks = _chunk_text_with_locations(text, "Raw Text")
        return chunks, "PARTIAL", 1

    chunks: List[Tuple[str, str]] = []

    def _traverse(obj: Any, path: str):
        if isinstance(obj, dict):
            for k, v in obj.items():
                cur_path = f"{path}.{k}" if path else k
                # Scan both the key and the string assignment
                if isinstance(v, (str, int, float, bool)) and v is not None:
                    chunks.append((f"{k}: {v}", f"JSON Path '{cur_path}'"))
                else:
                    _traverse(v, cur_path)
        elif isinstance(obj, list):
            for i, item in enumerate(obj):
                cur_path = f"{path}[{i}]"
                if isinstance(item, (str, int, float, bool)) and item is not None:
                    chunks.append((str(item), f"JSON Path '{cur_path}'"))
                else:
                    _traverse(item, cur_path)

    _traverse(data, "")
    if not chunks:
        # If empty structure, keep raw string
        chunks = [(text, "JSON Root")]

    return chunks, "SUCCESS", len(chunks)

def _extract_from_pdf(file_bytes: bytes, filename: str) -> Tuple[List[Tuple[str, str]], str, int, str]:
    """
    Extract readable text from PDF documents using pypdf.
    Handles multi-page PDFs, encrypted files, and flags image-only/scanned pages.
    """
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    except Exception as e:
        return [("", "PDF Corrupted")], "FAILED", 0, "unavailable"

    if reader.is_encrypted:
        return [("", "PDF Encrypted")], "FAILED", 0, "unavailable"

    total_pages = len(reader.pages)
    if total_pages == 0:
        return [("", "PDF Empty")], "FAILED", 0, "unavailable"

    chunks: List[Tuple[str, str]] = []
    extraction_status = "SUCCESS"
    ocr_status = "unavailable"
    has_empty_pages = False

    for page_idx, page in enumerate(reader.pages, start=1):
        page_text = ""
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""

        cleaned = page_text.strip()
        if cleaned:
            page_chunks = _chunk_text_with_locations(cleaned, f"PDF Page {page_idx}")
            chunks.extend(page_chunks)
        else:
            has_empty_pages = True
            # Attempt OCR on embedded page images if available
            images_found = getattr(page, "images", [])
            page_ocr_success = False
            if images_found and is_ocr_available():
                ocr_status = "executed"
                for img_obj in images_found:
                    try:
                        from PIL import Image
                        pil_img = Image.open(io.BytesIO(img_obj.data))
                        ocr_txt, o_stat = extract_text_from_image(pil_img)
                        if ocr_txt and ocr_txt.strip():
                            chunks.append((ocr_txt.strip(), f"PDF Page {page_idx} (OCR Image)"))
                            page_ocr_success = True
                    except Exception:
                        pass
            
            if not page_ocr_success:
                extraction_status = "PARTIAL"

    if not chunks:
        extraction_status = "PARTIAL"

    return chunks, extraction_status, total_pages, ocr_status

def _extract_from_docx(file_bytes: bytes, filename: str) -> Tuple[List[Tuple[str, str]], str, int]:
    """Extract paragraphs and table text from DOCX files using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        return [("", "DOCX Corrupted")], "FAILED", 0

    chunks: List[Tuple[str, str]] = []
    paragraph_count = 0

    # Extract paragraphs
    for p_idx, p in enumerate(doc.paragraphs, start=1):
        txt = p.text.strip()
        if txt:
            paragraph_count += 1
            chunks.extend(_chunk_text_with_locations(txt, f"DOCX Paragraph {p_idx}"))

    # Extract tables
    for t_idx, table in enumerate(doc.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_vals:
                row_str = " | ".join(row_vals)
                chunks.append((row_str, f"DOCX Table {t_idx}, Row {r_idx}"))

    return chunks, "SUCCESS", paragraph_count

def scan_uploaded_file(
    file_bytes: bytes,
    filename: str,
    request_id: Optional[str] = None
) -> FileScanResponse:
    """
    Main document security scanner: Validates, extracts, and inspects files
    for sensitive data, credentials, and prompt injection attacks.
    """
    req_id = request_id or f"file_{uuid.uuid4().hex[:8]}"
    now_iso = datetime.now(timezone.utc).isoformat()
    file_size = len(file_bytes)
    ext = _detect_file_extension(filename)

    # 1. Size Validation
    if file_size > settings.MAX_FILE_SIZE_BYTES:
        return FileScanResponse(
            request_id=req_id,
            action="BLOCK",
            risk_level="high",
            filename=filename,
            detected_file_type=ext or "unknown",
            file_size_bytes=file_size,
            extraction_status="OVERSIZED",
            ocr_status="skipped",
            total_pages_or_rows=0,
            findings=[
                Finding(
                    type="RULE_FILE_SIZE_EXCEEDED",
                    category="Policy Violation",
                    description=f"File size ({file_size / (1024*1024):.1f} MB) exceeds maximum gateway limit of {settings.MAX_FILE_SIZE_BYTES / (1024*1024):.0f} MB.",
                    severity="high",
                    source_location="File Upload Header",
                    origin="document_text"
                )
            ],
            finding_count=1,
            scanned_components=[],
            sanitized_preview=None,
            extracted_text_preview=None,
            recommendation="Reject file. Reduce file size or compress document before scanning.",
            explanation="File exceeded maximum allowed payload size.",
            scanned_at=now_iso
        )

    # 2. Extension Allowlist Check
    if ext not in settings.ALLOWED_DOCUMENT_EXTENSIONS:
        return FileScanResponse(
            request_id=req_id,
            action="BLOCK",
            risk_level="high",
            filename=filename,
            detected_file_type=ext or "unknown",
            file_size_bytes=file_size,
            extraction_status="UNSUPPORTED",
            ocr_status="skipped",
            total_pages_or_rows=0,
            findings=[
                Finding(
                    type="RULE_FILE_UNSUPPORTED_FORMAT",
                    category="Policy Violation",
                    description=f"File extension '{ext}' is not supported. Supported formats: {', '.join(settings.ALLOWED_DOCUMENT_EXTENSIONS)}",
                    severity="high",
                    source_location="Filename",
                    origin="document_text"
                )
            ],
            finding_count=1,
            scanned_components=[],
            sanitized_preview=None,
            extracted_text_preview=None,
            recommendation="Upload only supported document types (.pdf, .docx, .txt, .md, .csv, .json).",
            explanation=f"Unsupported file format '{ext}'.",
            scanned_at=now_iso
        )

    # 3. File Signature / Magic Bytes Validation
    is_valid_sig, sig_error = _validate_file_signature(file_bytes, ext)
    if not is_valid_sig:
        return FileScanResponse(
            request_id=req_id,
            action="BLOCK",
            risk_level="critical",
            filename=filename,
            detected_file_type=ext,
            file_size_bytes=file_size,
            extraction_status="FAILED",
            ocr_status="skipped",
            total_pages_or_rows=0,
            findings=[
                Finding(
                    type="RULE_FILE_MALFORMED_OR_CORRUPT",
                    category="Integrity Violation",
                    description=sig_error or "File signature validation failed.",
                    severity="critical",
                    source_location="File Header",
                    origin="document_text"
                )
            ],
            finding_count=1,
            scanned_components=[],
            sanitized_preview=None,
            extracted_text_preview=None,
            recommendation="Block corrupted or potentially masqueraded executable payload.",
            explanation=sig_error or "Malformed file header.",
            scanned_at=now_iso
        )

    # 4. Extract Text Content by Format
    chunks: List[Tuple[str, str]] = []
    extraction_status = "SUCCESS"
    ocr_status = "unavailable"
    total_units = 1

    try:
        if ext in (".txt", ".md"):
            chunks, extraction_status, total_units = _extract_from_txt_or_md(file_bytes, filename)
        elif ext == ".csv":
            chunks, extraction_status, total_units = _extract_from_csv(file_bytes, filename)
        elif ext == ".json":
            chunks, extraction_status, total_units = _extract_from_json(file_bytes, filename)
        elif ext == ".pdf":
            chunks, extraction_status, total_units, ocr_status = _extract_from_pdf(file_bytes, filename)
        elif ext == ".docx":
            chunks, extraction_status, total_units = _extract_from_docx(file_bytes, filename)
    except Exception as e:
        return FileScanResponse(
            request_id=req_id,
            action="BLOCK",
            risk_level="high",
            filename=filename,
            detected_file_type=ext,
            file_size_bytes=file_size,
            extraction_status="FAILED",
            ocr_status="failed",
            total_pages_or_rows=0,
            findings=[
                Finding(
                    type="RULE_FILE_PARSER_EXCEPTION",
                    category="Parser Failure",
                    description=f"Failed to parse document: {str(e)}",
                    severity="high",
                    source_location="File Parser",
                    origin="document_text"
                )
            ],
            finding_count=1,
            scanned_components=[],
            sanitized_preview=None,
            extracted_text_preview=None,
            recommendation="Verify document validity and retry without password protection.",
            explanation=f"Document parsing error: {str(e)}",
            scanned_at=now_iso
        )

    # If parser failed completely to extract readable stream
    if extraction_status == "FAILED":
        return FileScanResponse(
            request_id=req_id,
            action="BLOCK",
            risk_level="high",
            filename=filename,
            detected_file_type=ext,
            file_size_bytes=file_size,
            extraction_status="FAILED",
            ocr_status=ocr_status,
            total_pages_or_rows=0,
            findings=[
                Finding(
                    type="RULE_FILE_PARSE_FAILED",
                    category="Integrity Failure",
                    description="File could not be parsed or contains zero readable pages/records.",
                    severity="high",
                    source_location="Document Parser",
                    origin="document_text"
                )
            ],
            finding_count=1,
            scanned_components=[],
            sanitized_preview=None,
            extracted_text_preview=None,
            recommendation="Verify file integrity and upload an uncorrupted document.",
            explanation="Document parser failed to extract readable content.",
            scanned_at=now_iso
        )

    # 5. Security Scanning on Extracted Chunks
    policy = load_policies()
    aggregated_findings: List[Finding] = []
    scanned_locations = []
    all_redacted_texts = []
    all_raw_texts = []

    seen_finding_signatures = set()

    for chunk_text, location in chunks:
        if not chunk_text.strip():
            continue

        scanned_locations.append(location)
        all_raw_texts.append(chunk_text)

        # 5a. Sensitive data scan & redaction
        s_findings, redacted_chunk = scan_sensitive_data(chunk_text, policy)
        all_redacted_texts.append(redacted_chunk)

        for sf in s_findings:
            sig = (sf["type"], sf["match_preview"], location)
            if sig not in seen_finding_signatures:
                seen_finding_signatures.add(sig)
                aggregated_findings.append(Finding(
                    type=sf["type"],
                    category=sf["category"],
                    description=sf["description"],
                    severity=sf["severity"],
                    match_preview=sf.get("match_preview"),
                    start=sf.get("start"),
                    end=sf.get("end"),
                    source_location=location,
                    origin="ocr" if "OCR" in location else "document_text"
                ))

        # 5b. Prompt injection & heuristic scan
        inj_findings, inj_risk, _, _ = detect_prompt_injection(chunk_text, policy)
        for inf in inj_findings:
            sig = (inf["type"], inf["description"], location)
            if sig not in seen_finding_signatures:
                seen_finding_signatures.add(sig)
                aggregated_findings.append(Finding(
                    type=inf["type"],
                    category=inf["category"],
                    description=inf["description"],
                    severity=inf["severity"],
                    match_preview=inf.get("match_preview"),
                    source_location=location,
                    origin="ocr" if "OCR" in location else "document_text"
                ))

    # 6. Central Policy Evaluation
    sensitive_dict_list = [f.model_dump() for f in aggregated_findings if "Injection" not in f.category and "Jailbreak" not in f.category and "Prompt" not in f.category]
    injection_dict_list = [f.model_dump() for f in aggregated_findings if "Injection" in f.category or "Jailbreak" in f.category or "Prompt" in f.category]
    highest_inj_risk = "high" if any(f.get("severity") in ("high", "critical") for f in injection_dict_list) else "none"

    policy_verdict = evaluate_security_policy(
        sensitive_findings=sensitive_dict_list,
        injection_findings=injection_dict_list,
        injection_risk=highest_inj_risk,
        policy=policy
    )

    action = policy_verdict["action"]
    risk_level = policy_verdict["risk_level"]
    recommendation = policy_verdict["recommendation"]
    explanation = policy_verdict["explanation"]

    # If extraction was incomplete/partial and no findings were detected, mark risk as medium
    if extraction_status == "PARTIAL" and not aggregated_findings:
        if risk_level == "low":
            risk_level = "medium"
            recommendation = "Text extraction was incomplete (e.g. scanned image pages). File cannot be verified as completely clean."
            explanation = "Document contains unextracted or scanned elements without text."

    # Generate safe previews (maximum 160 chars)
    full_extracted_sample = " ".join(all_raw_texts)
    extracted_preview = (full_extracted_sample[:157] + "...") if len(full_extracted_sample) > 160 else full_extracted_sample
    
    full_redacted_sample = " ".join(all_redacted_texts)
    sanitized_preview = (full_redacted_sample[:157] + "...") if len(full_redacted_sample) > 160 else full_redacted_sample

    return FileScanResponse(
        request_id=req_id,
        action=action,
        risk_level=risk_level,
        filename=filename,
        detected_file_type=ext,
        file_size_bytes=file_size,
        extraction_status=extraction_status,
        ocr_status=ocr_status,
        total_pages_or_rows=total_units,
        findings=aggregated_findings,
        finding_count=len(aggregated_findings),
        scanned_components=scanned_locations[:15],
        sanitized_preview=sanitized_preview if sanitized_preview else None,
        extracted_text_preview=sanitized_preview if action != "ALLOW" else extracted_preview,
        recommendation=recommendation,
        explanation=explanation,
        scanned_at=now_iso
    )
